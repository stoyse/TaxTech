import streamlit as st
import streamlit.components.v1 as components
import json
from config import get_active_config, load_system_prompt, AI_PROVIDER
from openai import OpenAI
import PyPDF2
import docx
import io
import logging
from datetime import datetime

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s',
    handlers=[
        logging.FileHandler('doc2flow.log'),
        logging.StreamHandler()
    ]
)
logger = logging.getLogger(__name__)

# Initialize AI client
def get_ai_client():
    """Initialize the AI client based on the configured provider"""
    logger.info(f"Initializing AI client for provider: {AI_PROVIDER}")
    config = get_active_config()
    
    if not config["api_key"]:
        logger.error(f"No API key found for provider: {AI_PROVIDER}")
        return None
    
    logger.info(f"API key found (length: {len(config['api_key'])})")
    
    if AI_PROVIDER == "openai":
        client = OpenAI(api_key=config["api_key"])
        logger.info("OpenAI client initialized successfully")
        return client
    # Add Anthropic and Gemini support later
    logger.warning(f"Provider {AI_PROVIDER} not yet supported")
    return None

# Helper function to extract text from uploaded files
def extract_text_from_file(uploaded_file):
    """Extract text content from uploaded files"""
    logger.info(f"Extracting text from file: {uploaded_file.name}")
    try:
        file_type = uploaded_file.name.split('.')[-1].lower()
        logger.info(f"File type detected: {file_type}")
        
        if file_type == 'txt' or file_type == 'md':
            text = uploaded_file.read().decode('utf-8')
            logger.info(f"Extracted {len(text)} characters from {file_type} file")
            return text
        
        elif file_type == 'pdf':
            pdf_reader = PyPDF2.PdfReader(io.BytesIO(uploaded_file.read()))
            text = ""
            logger.info(f"PDF has {len(pdf_reader.pages)} pages")
            for i, page in enumerate(pdf_reader.pages):
                page_text = page.extract_text()
                logger.info(f"Page {i+1}: extracted {len(page_text)} characters")
                text += page_text + "\n"
            logger.info(f"Total extracted from PDF: {len(text)} characters")
            return text
        
        elif file_type == 'docx':
            doc = docx.Document(io.BytesIO(uploaded_file.read()))
            text = ""
            logger.info(f"DOCX has {len(doc.paragraphs)} paragraphs")
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            logger.info(f"Total extracted from DOCX: {len(text)} characters")
            return text
        
        else:
            logger.warning(f"Unsupported file type: {file_type}")
            return ""
    except Exception as e:
        logger.error(f"Error extracting text from {uploaded_file.name}: {str(e)}", exc_info=True)
        st.error(f"Error extracting text from {uploaded_file.name}: {str(e)}")
        return ""

# Function to generate BPMN using AI
def generate_bpmn_from_documents(uploaded_files):
    """Generate BPMN workflow from uploaded documents using AI"""
    logger.info("="*80)
    logger.info("STARTING BPMN GENERATION PROCESS")
    logger.info(f"Number of uploaded files: {len(uploaded_files)}")
    
    try:
        # Extract text from all uploaded files
        combined_text = ""
        for file in uploaded_files:
            logger.info(f"Processing file: {file.name}")
            text = extract_text_from_file(file)
            combined_text += f"\n\n=== Content from {file.name} ===\n{text}"
        
        logger.info(f"Total combined text length: {len(combined_text)} characters")
        
        if not combined_text.strip():
            logger.error("No text extracted from uploaded files")
            return None, "Could not extract text from the uploaded files."
        
        # Get AI configuration
        logger.info("Getting AI configuration...")
        config = get_active_config()
        
        # Get the token limit parameter (different providers use different names)
        token_limit_key = next((k for k in ["max_completion_tokens", "max_tokens", "max_output_tokens"] if k in config), None)
        token_limit = config.get(token_limit_key, "N/A") if token_limit_key else "N/A"
        logger.info(f"Config retrieved: model={config['model']}, temperature={config['temperature']}, {token_limit_key}={token_limit}")
        
        logger.info("Getting AI client...")
        client = get_ai_client()
        
        if not client:
            logger.error(f"Failed to initialize AI client for provider: {AI_PROVIDER}")
            return None, f"AI provider {AI_PROVIDER} is not yet supported or API key is missing."
        
        # Load system prompt
        logger.info("Loading system prompt...")
        system_prompt = load_system_prompt()
        logger.info(f"System prompt loaded: {len(system_prompt)} characters")
        
        # Create the prompt for BPMN generation
        text_preview = combined_text[:8000]
        logger.info(f"Using first {len(text_preview)} characters for API call")
        
        user_prompt = f"""Based on the following document content, analyze the business process described and generate a complete BPMN 2.0 XML workflow.

Document Content:
{text_preview}  

Please create a BPMN workflow that:
1. Identifies the main process steps
2. Includes start and end events
3. Adds decision points (gateways) where applicable
4. Uses clear, descriptive labels
5. Includes complete diagram interchange (DI) information for proper visual rendering

Return ONLY the BPMN XML without any additional explanation or markdown formatting."""
        
        logger.info("Calling OpenAI API...")
        logger.info(f"System prompt preview: {system_prompt[:200]}...")
        logger.info(f"User prompt preview: {user_prompt[:500]}...")
        
        # Call OpenAI API
        # Use max_completion_tokens for newer models (gpt-4o, etc.) or max_tokens for older models
        api_params = {
            "model": config["model"],
            "messages": [
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": user_prompt}
            ],
            "temperature": config["temperature"],
        }
        
        # Add the appropriate token limit parameter based on what's in config
        if "max_completion_tokens" in config:
            api_params["max_completion_tokens"] = config["max_completion_tokens"]
        elif "max_tokens" in config:
            api_params["max_tokens"] = config["max_tokens"]
        
        response = client.chat.completions.create(**api_params)
        
        logger.info("API call successful!")
        logger.info(f"Response ID: {response.id}")
        logger.info(f"Model used: {response.model}")
        logger.info(f"Tokens used - prompt: {response.usage.prompt_tokens}, completion: {response.usage.completion_tokens}, total: {response.usage.total_tokens}")
        
        bpmn_xml = response.choices[0].message.content.strip()
        logger.info(f"Raw response length: {len(bpmn_xml)} characters")
        logger.info(f"Raw response preview (first 500 chars):\n{bpmn_xml[:500]}")
        
        # Clean up the response (remove markdown code blocks if present)
        if bpmn_xml.startswith("```xml"):
            logger.info("Removing ```xml markdown wrapper")
            bpmn_xml = bpmn_xml[6:]
        if bpmn_xml.startswith("```"):
            logger.info("Removing ``` markdown wrapper")
            bpmn_xml = bpmn_xml[3:]
        if bpmn_xml.endswith("```"):
            logger.info("Removing trailing ``` markdown wrapper")
            bpmn_xml = bpmn_xml[:-3]
        bpmn_xml = bpmn_xml.strip()
        
        logger.info(f"Cleaned BPMN XML length: {len(bpmn_xml)} characters")
        logger.info(f"BPMN XML preview (first 500 chars):\n{bpmn_xml[:500]}")
        logger.info("BPMN GENERATION COMPLETED SUCCESSFULLY")
        logger.info("="*80)
        
        return bpmn_xml, "Successfully generated BPMN workflow!"
        
    except Exception as e:
        logger.error(f"ERROR generating BPMN: {str(e)}", exc_info=True)
        logger.info("="*80)
        return None, f"Error generating BPMN: {str(e)}"

# Page configuration
st.set_page_config(
    page_title="Doc2Flow",
    page_icon="📄",
    layout="wide"
)

# Main Title
st.title("📄 Doc2Flow")
st.markdown("### Transform your documents into clear BPMN workflows")

# Sidebar for configuration and logs
with st.sidebar:
    st.header("⚙️ Configuration")
    st.info(f"**AI Provider:** {AI_PROVIDER}")
    config = get_active_config()
    st.info(f"**Model:** {config['model']}")
    
    st.markdown("---")
    
    # Log viewer
    st.header("📋 Debug Logs")
    show_logs = st.checkbox("Show real-time logs", value=False)
    
    if show_logs:
        try:
            with open('doc2flow.log', 'r') as f:
                logs = f.readlines()
                # Show last 50 lines
                recent_logs = ''.join(logs[-50:])
                st.text_area("Recent Logs", recent_logs, height=400)
        except FileNotFoundError:
            st.info("No logs yet. Start using the app to generate logs.")

# Create main layout with two columns
left_col, right_col = st.columns([1, 1])

# Left column - Chat Interface
with left_col:
    st.markdown("#### 💬 Chat & Upload")
    
    # Initialize chat history
    if "messages" not in st.session_state:
        st.session_state.messages = []
    
    # File uploader in the chat area
    uploaded_files = st.file_uploader(
        "📎 Upload documents (PDF, DOCX, TXT, MD)",
        accept_multiple_files=True,
        type=['pdf', 'docx', 'txt', 'md'],
        help="Upload documents to analyze and generate BPMN"
    )
    
    if uploaded_files and "last_uploaded" not in st.session_state:
        st.session_state.last_uploaded = [f.name for f in uploaded_files]
        file_list = ", ".join([f.name for f in uploaded_files])
        st.session_state.messages.append({
            "role": "assistant", 
            "content": f"✅ I've received {len(uploaded_files)} file(s): {file_list}. Would you like me to generate a BPMN workflow from these documents?"
        })
    
    # Chat container with fixed height
    chat_container = st.container(height=400)
    
    with chat_container:
        # Display chat messages
        for message in st.session_state.messages:
            with st.chat_message(message["role"]):
                st.markdown(message["content"])
    
    # Chat input
    if prompt := st.chat_input("Ask me to generate, modify, or explain the BPMN workflow..."):
        logger.info(f"[CHAT] User prompt received: {prompt}")
        # Add user message to chat history
        st.session_state.messages.append({"role": "user", "content": prompt})
        
        # Process user request
        with st.spinner("Analyzing documents and generating BPMN..."):
            response = ""
            if any(word in prompt.lower() for word in ["generate", "create", "make", "build"]):
                logger.info("[CHAT] Generation keywords detected in prompt")
                # Check if files are uploaded
                if uploaded_files:
                    logger.info(f"[CHAT] Files are uploaded: {[f.name for f in uploaded_files]}")
                    # Generate BPMN workflow from uploaded documents using AI
                    bpmn_xml, message = generate_bpmn_from_documents(uploaded_files)
                    
                    if bpmn_xml:
                        logger.info("[CHAT] BPMN generated successfully, storing in session state")
                        st.session_state.generated_workflow = bpmn_xml
                        response = f"✅ {message} I've analyzed your documents and created a custom BPMN workflow. You can see it on the right side and edit it as needed."
                    else:
                        logger.error(f"[CHAT] BPMN generation failed: {message}")
                        response = f"❌ {message}"
                else:
                    logger.warning("[CHAT] No files uploaded")
                    response = "⚠️ Please upload some documents first before I can generate a BPMN workflow."
                    # Fallback to default workflow
                    st.session_state.generated_workflow = '''<?xml version="1.0" encoding="UTF-8"?>
<bpmn:definitions xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance" xmlns:bpmn="http://www.omg.org/spec/BPMN/20100524/MODEL" xmlns:bpmndi="http://www.omg.org/spec/BPMN/20100524/DI" xmlns:dc="http://www.omg.org/spec/DD/20100524/DC" xmlns:di="http://www.omg.org/spec/DD/20100524/DI" id="Definitions_1" targetNamespace="http://bpmn.io/schema/bpmn">
  <bpmn:process id="Process_1" isExecutable="false">
    <bpmn:startEvent id="start" name="Document Received">
      <bpmn:outgoing>flow1</bpmn:outgoing>
    </bpmn:startEvent>
    <bpmn:exclusiveGateway id="validate" name="Validate Document">
      <bpmn:incoming>flow1</bpmn:incoming>
      <bpmn:outgoing>flow2</bpmn:outgoing>
    </bpmn:exclusiveGateway>
    <bpmn:task id="process" name="Process Content">
      <bpmn:incoming>flow2</bpmn:incoming>
      <bpmn:outgoing>flow3</bpmn:outgoing>
    </bpmn:task>
    <bpmn:task id="extract" name="Extract Information">
      <bpmn:incoming>flow3</bpmn:incoming>
      <bpmn:outgoing>flow4</bpmn:outgoing>
    </bpmn:task>
    <bpmn:exclusiveGateway id="review" name="Manual Review">
      <bpmn:incoming>flow4</bpmn:incoming>
      <bpmn:outgoing>flow5</bpmn:outgoing>
    </bpmn:exclusiveGateway>
    <bpmn:task id="approve" name="Approve Process">
      <bpmn:incoming>flow5</bpmn:incoming>
      <bpmn:outgoing>flow6</bpmn:outgoing>
    </bpmn:task>
    <bpmn:endEvent id="end" name="Process Complete">
      <bpmn:incoming>flow6</bpmn:incoming>
    </bpmn:endEvent>
    <bpmn:sequenceFlow id="flow1" sourceRef="start" targetRef="validate"/>
    <bpmn:sequenceFlow id="flow2" sourceRef="validate" targetRef="process"/>
    <bpmn:sequenceFlow id="flow3" sourceRef="process" targetRef="extract"/>
    <bpmn:sequenceFlow id="flow4" sourceRef="extract" targetRef="review"/>
    <bpmn:sequenceFlow id="flow5" sourceRef="review" targetRef="approve"/>
    <bpmn:sequenceFlow id="flow6" sourceRef="approve" targetRef="end"/>
  </bpmn:process>
  <bpmndi:BPMNDiagram id="BPMNDiagram_1">
    <bpmndi:BPMNPlane id="BPMNPlane_1" bpmnElement="Process_1">
      <bpmndi:BPMNShape id="start_di" bpmnElement="start">
        <dc:Bounds x="152" y="102" width="36" height="36"/>
        <bpmndi:BPMNLabel>
          <dc:Bounds x="144" y="145" width="52" height="27"/>
        </bpmndi:BPMNLabel>
      </bpmndi:BPMNShape>
      <bpmndi:BPMNShape id="validate_di" bpmnElement="validate" isMarkerVisible="true">
        <dc:Bounds x="245" y="95" width="50" height="50"/>
        <bpmndi:BPMNLabel>
          <dc:Bounds x="239" y="152" width="62" height="27"/>
        </bpmndi:BPMNLabel>
      </bpmndi:BPMNShape>
      <bpmndi:BPMNShape id="process_di" bpmnElement="process">
        <dc:Bounds x="360" y="80" width="100" height="80"/>
      </bpmndi:BPMNShape>
      <bpmndi:BPMNShape id="extract_di" bpmnElement="extract">
        <dc:Bounds x="520" y="80" width="100" height="80"/>
      </bpmndi:BPMNShape>
      <bpmndi:BPMNShape id="review_di" bpmnElement="review" isMarkerVisible="true">
        <dc:Bounds x="680" y="95" width="50" height="50"/>
        <bpmndi:BPMNLabel>
          <dc:Bounds x="665" y="152" width="80" height="14"/>
        </bpmndi:BPMNLabel>
      </bpmndi:BPMNShape>
      <bpmndi:BPMNShape id="approve_di" bpmnElement="approve">
        <dc:Bounds x="790" y="80" width="100" height="80"/>
      </bpmndi:BPMNShape>
      <bpmndi:BPMNShape id="end_di" bpmnElement="end">
        <dc:Bounds x="952" y="102" width="36" height="36"/>
        <bpmndi:BPMNLabel>
          <dc:Bounds x="947" y="145" width="46" height="27"/>
        </bpmndi:BPMNLabel>
      </bpmndi:BPMNShape>
      <bpmndi:BPMNEdge id="flow1_di" bpmnElement="flow1">
        <di:waypoint x="188" y="120"/>
        <di:waypoint x="245" y="120"/>
      </bpmndi:BPMNEdge>
      <bpmndi:BPMNEdge id="flow2_di" bpmnElement="flow2">
        <di:waypoint x="295" y="120"/>
        <di:waypoint x="360" y="120"/>
      </bpmndi:BPMNEdge>
      <bpmndi:BPMNEdge id="flow3_di" bpmnElement="flow3">
        <di:waypoint x="460" y="120"/>
        <di:waypoint x="520" y="120"/>
      </bpmndi:BPMNEdge>
      <bpmndi:BPMNEdge id="flow4_di" bpmnElement="flow4">
        <di:waypoint x="620" y="120"/>
        <di:waypoint x="680" y="120"/>
      </bpmndi:BPMNEdge>
      <bpmndi:BPMNEdge id="flow5_di" bpmnElement="flow5">
        <di:waypoint x="730" y="120"/>
        <di:waypoint x="790" y="120"/>
      </bpmndi:BPMNEdge>
      <bpmndi:BPMNEdge id="flow6_di" bpmnElement="flow6">
        <di:waypoint x="890" y="120"/>
        <di:waypoint x="952" y="120"/>
      </bpmndi:BPMNEdge>
    </bpmndi:BPMNPlane>
  </bpmndi:BPMNDiagram>
</bpmn:definitions>'''
                response = "✅ BPMN workflow generated successfully! You can now see it on the right side and edit it as needed."
                
            elif any(word in prompt.lower() for word in ["modify", "change", "update", "edit", "add", "remove"]):
                response = "I can help you modify the workflow. You can directly edit the BPMN diagram on the right side by:\n- Dragging elements to reposition them\n- Clicking elements to edit their properties\n- Using the palette on the left to add new elements\n- Selecting elements and pressing Delete to remove them\n\nWhat specific changes would you like to make?"
                
            elif any(word in prompt.lower() for word in ["explain", "what", "how", "why"]):
                response = "This BPMN workflow shows the process flow with:\n- **Start Events** (circles): Entry points\n- **Tasks** (rectangles): Activities to perform\n- **Gateways** (diamonds): Decision points\n- **End Events** (bold circles): Process completion\n- **Sequence Flows** (arrows): Process flow direction\n\nYou can edit the diagram directly on the right side. What would you like to know more about?"
                
            elif any(word in prompt.lower() for word in ["download", "export", "save"]):
                response = "📥 You can download the workflow as an SVG file using the 'Download BPMN as SVG' button below the diagram. The SVG format maintains quality at any size and can be imported into other tools."
                
            else:
                response = "I can help you with:\n- 🔨 **Generate** BPMN workflows from documents\n- ✏️ **Modify** existing workflows\n- 📖 **Explain** BPMN elements and processes\n- 📥 **Download** workflows as SVG\n\nJust ask me what you'd like to do!"
            
            # Add assistant response to chat history
            st.session_state.messages.append({"role": "assistant", "content": response})
        
        # Rerun to display new messages
        st.rerun()

# Right column - Workflow Viewer
with right_col:
    st.markdown("#### BPMN Workflow Viewer")

    # Check if we have a generated workflow from the dashboard
    if 'generated_workflow' in st.session_state:
        st.info("📄 Displaying workflow generated from your uploaded documents")
        default_workflow = st.session_state.generated_workflow
    else:
        default_workflow = '''<?xml version="1.0" encoding="UTF-8"?>
<bpmn:definitions xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance" xmlns:bpmn="http://www.omg.org/spec/BPMN/20100524/MODEL" xmlns:bpmndi="http://www.omg.org/spec/BPMN/20100524/DI" xmlns:dc="http://www.omg.org/spec/DD/20100524/DC" xmlns:di="http://www.omg.org/spec/DD/20100524/DI" id="Definitions_1" targetNamespace="http://bpmn.io/schema/bpmn">
  <bpmn:process id="Process_1" isExecutable="false">
    <bpmn:startEvent id="start" name="Receive Loan Request">
      <bpmn:outgoing>flow1</bpmn:outgoing>
    </bpmn:startEvent>
    <bpmn:exclusiveGateway id="gateway1" name="Check Eligibility">
      <bpmn:incoming>flow1</bpmn:incoming>
      <bpmn:outgoing>flow2</bpmn:outgoing>
    </bpmn:exclusiveGateway>
    <bpmn:exclusiveGateway id="gateway2" name="Evaluate Risk">
      <bpmn:incoming>flow2</bpmn:incoming>
      <bpmn:outgoing>flow3</bpmn:outgoing>
      <bpmn:outgoing>flow4</bpmn:outgoing>
    </bpmn:exclusiveGateway>
    <bpmn:task id="approve" name="Approve Loan">
      <bpmn:incoming>flow3</bpmn:incoming>
      <bpmn:outgoing>flow5</bpmn:outgoing>
    </bpmn:task>
    <bpmn:task id="reject" name="Reject Application">
      <bpmn:incoming>flow4</bpmn:incoming>
      <bpmn:outgoing>flow6</bpmn:outgoing>
    </bpmn:task>
    <bpmn:endEvent id="end" name="Send Response">
      <bpmn:incoming>flow5</bpmn:incoming>
      <bpmn:incoming>flow6</bpmn:incoming>
    </bpmn:endEvent>
    <bpmn:sequenceFlow id="flow1" sourceRef="start" targetRef="gateway1"/>
    <bpmn:sequenceFlow id="flow2" sourceRef="gateway1" targetRef="gateway2"/>
    <bpmn:sequenceFlow id="flow3" sourceRef="gateway2" targetRef="approve"/>
    <bpmn:sequenceFlow id="flow4" sourceRef="gateway2" targetRef="reject"/>
    <bpmn:sequenceFlow id="flow5" sourceRef="approve" targetRef="end"/>
    <bpmn:sequenceFlow id="flow6" sourceRef="reject" targetRef="end"/>
  </bpmn:process>
  <bpmndi:BPMNDiagram id="BPMNDiagram_1">
    <bpmndi:BPMNPlane id="BPMNPlane_1" bpmnElement="Process_1">
      <bpmndi:BPMNShape id="start_di" bpmnElement="start">
        <dc:Bounds x="152" y="102" width="36" height="36"/>
        <bpmndi:BPMNLabel>
          <dc:Bounds x="139" y="145" width="62" height="27"/>
        </bpmndi:BPMNLabel>
      </bpmndi:BPMNShape>
      <bpmndi:BPMNShape id="gateway1_di" bpmnElement="gateway1" isMarkerVisible="true">
        <dc:Bounds x="245" y="95" width="50" height="50"/>
        <bpmndi:BPMNLabel>
          <dc:Bounds x="244" y="152" width="52" height="27"/>
        </bpmndi:BPMNLabel>
      </bpmndi:BPMNShape>
      <bpmndi:BPMNShape id="gateway2_di" bpmnElement="gateway2" isMarkerVisible="true">
        <dc:Bounds x="355" y="95" width="50" height="50"/>
        <bpmndi:BPMNLabel>
          <dc:Bounds x="345" y="152" width="70" height="14"/>
        </bpmndi:BPMNLabel>
      </bpmndi:BPMNShape>
      <bpmndi:BPMNShape id="approve_di" bpmnElement="approve">
        <dc:Bounds x="470" y="80" width="100" height="80"/>
      </bpmndi:BPMNShape>
      <bpmndi:BPMNShape id="reject_di" bpmnElement="reject">
        <dc:Bounds x="470" y="200" width="100" height="80"/>
      </bpmndi:BPMNShape>
      <bpmndi:BPMNShape id="end_di" bpmnElement="end">
        <dc:Bounds x="632" y="162" width="36" height="36"/>
        <bpmndi:BPMNLabel>
          <dc:Bounds x="611" y="205" width="78" height="14"/>
        </bpmndi:BPMNLabel>
      </bpmndi:BPMNShape>
      <bpmndi:BPMNEdge id="flow1_di" bpmnElement="flow1">
        <di:waypoint x="188" y="120"/>
        <di:waypoint x="245" y="120"/>
      </bpmndi:BPMNEdge>
      <bpmndi:BPMNEdge id="flow2_di" bpmnElement="flow2">
        <di:waypoint x="295" y="120"/>
        <di:waypoint x="355" y="120"/>
      </bpmndi:BPMNEdge>
      <bpmndi:BPMNEdge id="flow3_di" bpmnElement="flow3">
        <di:waypoint x="405" y="120"/>
        <di:waypoint x="470" y="120"/>
      </bpmndi:BPMNEdge>
      <bpmndi:BPMNEdge id="flow4_di" bpmnElement="flow4">
        <di:waypoint x="380" y="145"/>
        <di:waypoint x="380" y="240"/>
        <di:waypoint x="470" y="240"/>
      </bpmndi:BPMNEdge>
      <bpmndi:BPMNEdge id="flow5_di" bpmnElement="flow5">
        <di:waypoint x="570" y="120"/>
        <di:waypoint x="601" y="120"/>
        <di:waypoint x="601" y="180"/>
        <di:waypoint x="632" y="180"/>
      </bpmndi:BPMNEdge>
      <bpmndi:BPMNEdge id="flow6_di" bpmnElement="flow6">
        <di:waypoint x="570" y="240"/>
        <di:waypoint x="601" y="240"/>
        <di:waypoint x="601" y="180"/>
        <di:waypoint x="632" y="180"/>
      </bpmndi:BPMNEdge>
    </bpmndi:BPMNPlane>
  </bpmndi:BPMNDiagram>
</bpmn:definitions>'''

    # Control buttons
    btn_col1, btn_col2 = st.columns([1, 1])
    
    with btn_col1:
        if st.button("🚀 Generate Workflow", type="primary", use_container_width=True):
            logger.info("[BUTTON] Generate Workflow button clicked")
            if uploaded_files:
                logger.info(f"[BUTTON] Files detected: {[f.name for f in uploaded_files]}")
                with st.spinner("Analyzing documents and generating BPMN..."):
                    bpmn_xml, message = generate_bpmn_from_documents(uploaded_files)
                    if bpmn_xml:
                        logger.info("[BUTTON] BPMN generated successfully, storing in session state")
                        logger.info(f"[BUTTON] Stored workflow length: {len(bpmn_xml)} characters")
                        st.session_state.generated_workflow = bpmn_xml
                        st.success(message)
                        generate_workflow = True
                        logger.info("[BUTTON] Triggering rerun to display new workflow")
                        st.rerun()
                    else:
                        logger.error(f"[BUTTON] BPMN generation failed: {message}")
                        st.error(message)
                        generate_workflow = False
            else:
                logger.warning("[BUTTON] No files uploaded")
                st.warning("⚠️ Please upload documents first!")
                generate_workflow = False
        else:
            generate_workflow = False
    
    with btn_col2:
        if st.button("🔄 Load Sample", use_container_width=True):
            st.session_state.generated_workflow = '''<?xml version="1.0" encoding="UTF-8"?>
<bpmn:definitions xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance" xmlns:bpmn="http://www.omg.org/spec/BPMN/20100524/MODEL" xmlns:bpmndi="http://www.omg.org/spec/BPMN/20100524/DI" xmlns:dc="http://www.omg.org/spec/DD/20100524/DC" xmlns:di="http://www.omg.org/spec/DD/20100524/DI" id="Definitions_1" targetNamespace="http://bpmn.io/schema/bpmn">
  <bpmn:process id="Process_1" isExecutable="false">
    <bpmn:startEvent id="start" name="Start Process">
      <bpmn:outgoing>flow1</bpmn:outgoing>
    </bpmn:startEvent>
    <bpmn:task id="input" name="Receive Input">
      <bpmn:incoming>flow1</bpmn:incoming>
      <bpmn:outgoing>flow2</bpmn:outgoing>
    </bpmn:task>
    <bpmn:exclusiveGateway id="validate" name="Validate Data">
      <bpmn:incoming>flow2</bpmn:incoming>
      <bpmn:outgoing>flow3</bpmn:outgoing>
    </bpmn:exclusiveGateway>
    <bpmn:task id="process" name="Process Request">
      <bpmn:incoming>flow3</bpmn:incoming>
      <bpmn:outgoing>flow4</bpmn:outgoing>
    </bpmn:task>
    <bpmn:endEvent id="end" name="End Process">
      <bpmn:incoming>flow4</bpmn:incoming>
    </bpmn:endEvent>
    <bpmn:sequenceFlow id="flow1" sourceRef="start" targetRef="input"/>
    <bpmn:sequenceFlow id="flow2" sourceRef="input" targetRef="validate"/>
    <bpmn:sequenceFlow id="flow3" sourceRef="validate" targetRef="process"/>
    <bpmn:sequenceFlow id="flow4" sourceRef="process" targetRef="end"/>
  </bpmn:process>
  <bpmndi:BPMNDiagram id="BPMNDiagram_1">
    <bpmndi:BPMNPlane id="BPMNPlane_1" bpmnElement="Process_1">
      <bpmndi:BPMNShape id="start_di" bpmnElement="start">
        <dc:Bounds x="152" y="102" width="36" height="36"/>
        <bpmndi:BPMNLabel>
          <dc:Bounds x="137" y="145" width="66" height="14"/>
        </bpmndi:BPMNLabel>
      </bpmndi:BPMNShape>
      <bpmndi:BPMNShape id="input_di" bpmnElement="input">
        <dc:Bounds x="250" y="80" width="100" height="80"/>
      </bpmndi:BPMNShape>
      <bpmndi:BPMNShape id="validate_di" bpmnElement="validate" isMarkerVisible="true">
        <dc:Bounds x="415" y="95" width="50" height="50"/>
        <bpmndi:BPMNLabel>
          <dc:Bounds x="400" y="152" width="80" height="14"/>
        </bpmndi:BPMNLabel>
      </bpmndi:BPMNShape>
      <bpmndi:BPMNShape id="process_di" bpmnElement="process">
        <dc:Bounds x="530" y="80" width="100" height="80"/>
      </bpmndi:BPMNShape>
      <bpmndi:BPMNShape id="end_di" bpmnElement="end">
        <dc:Bounds x="692" y="102" width="36" height="36"/>
        <bpmndi:BPMNLabel>
          <dc:Bounds x="677" y="145" width="66" height="14"/>
        </bpmndi:BPMNLabel>
      </bpmndi:BPMNShape>
      <bpmndi:BPMNEdge id="flow1_di" bpmnElement="flow1">
        <di:waypoint x="188" y="120"/>
        <di:waypoint x="250" y="120"/>
      </bpmndi:BPMNEdge>
      <bpmndi:BPMNEdge id="flow2_di" bpmnElement="flow2">
        <di:waypoint x="350" y="120"/>
        <di:waypoint x="415" y="120"/>
      </bpmndi:BPMNEdge>
      <bpmndi:BPMNEdge id="flow3_di" bpmnElement="flow3">
        <di:waypoint x="465" y="120"/>
        <di:waypoint x="530" y="120"/>
      </bpmndi:BPMNEdge>
      <bpmndi:BPMNEdge id="flow4_di" bpmnElement="flow4">
        <di:waypoint x="630" y="120"/>
        <di:waypoint x="692" y="120"/>
      </bpmndi:BPMNEdge>
    </bpmndi:BPMNPlane>
  </bpmndi:BPMNDiagram>
</bpmn:definitions>'''
            st.rerun()
    
    # Toggle for XML editor
    show_xml = st.toggle("📝 Show/Edit BPMN XML", value=False)
    
    # Use generated workflow if available, otherwise use default
    current_workflow = st.session_state.get('generated_workflow', default_workflow)
    logger.info(f"[VIEWER] Using workflow - length: {len(current_workflow)} characters")
    logger.info(f"[VIEWER] Workflow preview (first 200 chars): {current_workflow[:200]}...")
    
    # Workflow XML input (hidden behind toggle)
    if show_xml:
        bpmn_xml = st.text_area("BPMN XML", height=150, value=current_workflow)
    else:
        bpmn_xml = current_workflow

    # Display workflow if generated or sample loaded
    if generate_workflow or 'generated_workflow' in st.session_state:
        logger.info("[VIEWER] Rendering BPMN diagram...")
        logger.info(f"[VIEWER] BPMN XML to render - length: {len(bpmn_xml)}")
        logger.info(f"[VIEWER] BPMN XML preview (first 300 chars): {bpmn_xml[:300]}...")
        try:
            # Create the bpmn-js viewer HTML component
            # Replace backticks in XML to avoid breaking JavaScript template literal
            bpmn_xml_safe = bpmn_xml.replace('`', '\\`').replace('${', '\\${')
            logger.info("[VIEWER] BPMN XML escaped for JavaScript")
            
            bpmn_viewer_html = f"""
            <!DOCTYPE html>
            <html>
            <head>
                <link rel="stylesheet" href="https://unpkg.com/bpmn-js@17.2.0/dist/assets/bpmn-js.css">
                <link rel="stylesheet" href="https://unpkg.com/bpmn-js@17.2.0/dist/assets/diagram-js.css">
                <link rel="stylesheet" href="https://unpkg.com/bpmn-js@17.2.0/dist/assets/bpmn-font/css/bpmn-embedded.css">
                <style>
                    body {{
                        margin: 0;
                        padding: 0;
                        font-family: Arial, sans-serif;
                    }}
                    #canvas {{
                        height: 400px;
                        width: 100%;
                        border: 1px solid #e0e0e0;
                        border-radius: 4px;
                        background-color: white;
                    }}
                    .download-btn {{
                        margin: 10px 0;
                        padding: 8px 16px;
                        background-color: #0066cc;
                        color: white;
                        border: none;
                        border-radius: 4px;
                        cursor: pointer;
                        font-size: 14px;
                    }}
                    .download-btn:hover {{
                        background-color: #0052a3;
                    }}
                </style>
            </head>
            <body>
                <div id="canvas"></div>
                <button class="download-btn" onclick="downloadDiagram()">📥 Download BPMN as SVG</button>
                
                <script src="https://unpkg.com/bpmn-js@17.2.0/dist/bpmn-modeler.development.js"></script>
                <script>
                    var modeler = new BpmnJS({{
                        container: '#canvas',
                        keyboard: {{
                            bindTo: document
                        }}
                    }});
                    
                    var bpmnXML = `{bpmn_xml_safe}`;
                    
                    modeler.importXML(bpmnXML).then(function(result) {{
                        const canvas = modeler.get('canvas');
                        canvas.zoom('fit-viewport');
                    }}).catch(function(err) {{
                        console.error('Error rendering BPMN diagram:', err);
                    }});
                    
                    function downloadDiagram() {{
                        modeler.saveSVG().then(function(result) {{
                            var svg = result.svg;
                            var blob = new Blob([svg], {{ type: 'image/svg+xml' }});
                            var url = URL.createObjectURL(blob);
                            var a = document.createElement('a');
                            a.href = url;
                            a.download = 'workflow.svg';
                            document.body.appendChild(a);
                            a.click();
                            document.body.removeChild(a);
                            URL.revokeObjectURL(url);
                        }});
                    }}
                </script>
            </body>
            </html>
            """
            
            # Display the BPMN viewer
            components.html(bpmn_viewer_html, height=500)
                
        except Exception as e:
            st.error(f"Error rendering BPMN: {e}")
    else:
        st.info("👆 Upload documents and generate BPMN, or load a sample workflow to get started")